"""Augment a built EDM2-vs-flow comparison deck with rotating-3D-render
companion slides (docs/vivit_conditioning_plan.md Phase 5; team-lead request,
2026-09-13).

MUST run with the FLOW repo's interpreter plus python-pptx (same environment
as tools/build_flow_deck.sh):

    uv run --project D:/Work/GrowthNet_gamailab/FlowMatchingGrowthNet --with python-pptx python \\
        tools/augment_deck_3d.py \\
        --pptx <built deck.pptx> \\
        --old-analysis-dir <flow analysis dir> --new-analysis-dir <EDM2 analysis dir> \\
        --old-label "..." --new-label "..." \\
        --out <augmented deck.pptx>

Imports tumor_flow.analysis.build_comparison_deck.shared_scans() to get the
EXACT same ordered (patient_id, scan_id) identity and eyebrow-title text the
deck builder itself used for its per-scan slides, so a slide is matched by an
exact string match against text the deck builder wrote -- not a
re-derived/guessed pattern that could drift from it.

For every per-scan slide found this way, inserts a new companion slide
immediately after it with both models' rotating 3D GIFs (render_3d.gif, from
tools/render_edm2_3d.py / the flow repo's own pyvista_renderer.py) side by
side -- old/flow left, new/EDM2 right -- with labels and per-scan consensus-
Dice captions taken from each side's own viz_index.json metrics block (no
extra file reads needed). GIFs are embedded as plain pictures: current
desktop PowerPoint plays an embedded animated GIF during slideshow without
any special OLE wrapping. The flow repo's own render_3d.gif files run ~4.3MB,
so every GIF is downsampled (frame stride, then resolution) under
--max-gif-bytes (default 3MB) before embedding, working on a temporary copy;
the originals on disk are never modified.
"""

import argparse
import json
import os
import shutil
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from tumor_flow.analysis.build_comparison_deck import shared_scans

TEAL = RGBColor(0x0F, 0x76, 0x6E)
DARK = RGBColor(0x15, 0x22, 0x2B)

#----------------------------------------------------------------------------

def read_json(path):
    with open(path) as f:
        return json.load(f)


def gif_path(analysis_dir, patient_id, scan_id, name='render_3d.gif'):
    return os.path.join(analysis_dir, 'viz', f'patient{patient_id}', f'scan{scan_id}', name)

#----------------------------------------------------------------------------
# GIF downsampling: PIL can read/write animated GIFs frame-by-frame without
# pyvista/vtk, so this needs only python-pptx's own Pillow dependency.

def read_gif_frames(path):
    img = Image.open(path)
    frames, durations = [], []
    try:
        while True:
            frames.append(img.convert('RGB').copy())
            durations.append(img.info.get('duration', 100))
            img.seek(img.tell() + 1)
    except EOFError:
        pass
    return frames, durations


def write_gif(frames, durations, path):
    frames[0].save(
        path, save_all=True, append_images=frames[1:], duration=durations, loop=0, optimize=True)


def downsample_gif_under(src_path, dst_path, max_bytes):
    """Copy src_path to dst_path, shrinking (frame stride, then resolution)
    until under max_bytes or the most aggressive setting has been tried."""
    if os.path.getsize(src_path) <= max_bytes:
        shutil.copyfile(src_path, dst_path)
        return dst_path
    frames, durations = read_gif_frames(src_path)
    for stride in (2, 3, 4, 6, 8):
        for scale in (1.0, 0.75, 0.6, 0.5, 0.4, 0.3):
            sel_frames = frames[::stride]
            sel_durations = [d * stride for d in durations[::stride]]
            if scale != 1.0:
                w, h = frames[0].size
                new_size = (max(1, int(w * scale)), max(1, int(h * scale)))
                out_frames = [f.resize(new_size, Image.LANCZOS) for f in sel_frames]
            else:
                out_frames = sel_frames
            write_gif(out_frames, sel_durations, dst_path)
            if os.path.getsize(dst_path) <= max_bytes:
                return dst_path
    return dst_path  # most aggressive attempt already written to dst_path; accept it.

#----------------------------------------------------------------------------
# Slide identity/reordering.

def eyebrow_for(scan):
    # Byte-for-byte the same format string as
    # tumor_flow.analysis.build_comparison_deck._add_scan_slide's eyebrow.
    return f'PATIENT {scan.patient_id} - {scan.category.upper()} - +{scan.delta_days} d'


def slide_title_text(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.strip().splitlines()[0]
    return None


def find_blank_layout(prs):
    return next(layout for layout in prs.slide_layouts if layout.name == 'Blank')


def move_slide(prs, old_index, new_index):
    """Standard python-pptx slide-reordering recipe: python-pptx has no public
    API to insert a slide at a position, only to append and then reorder the
    presentation's own <p:sldIdLst> element."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])

#----------------------------------------------------------------------------

def add_companion_slide(prs, layout, scan, old_gif, new_gif, old_label, new_label):
    slide = prs.slides.add_slide(layout)

    eyebrow_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.4))
    eyebrow_box.text_frame.text = eyebrow_for(scan) + ' - 3D ROTATION'
    erun = eyebrow_box.text_frame.paragraphs[0].runs[0]
    erun.font.name, erun.font.size, erun.font.color.rgb = 'Consolas', Pt(12), TEAL

    old_dice = (scan.old_scan.get('metrics') or {}).get('consensus_dice')
    new_dice = (scan.new_scan.get('metrics') or {}).get('consensus_dice')

    col_width = Inches(6.0)
    col_gap = Inches(0.33)
    for col, (label, gif, dice) in enumerate((
        (old_label, old_gif, old_dice),
        (new_label, new_gif, new_dice),
    )):
        left = Inches(0.4) + col * (col_width + col_gap)
        caption = label if dice is None else f'{label}  (consensus Dice {dice:.3f})'
        label_box = slide.shapes.add_textbox(left, Inches(0.85), col_width, Inches(0.35))
        label_box.text_frame.text = caption
        lrun = label_box.text_frame.paragraphs[0].runs[0]
        lrun.font.name, lrun.font.size, lrun.font.color.rgb = 'Consolas', Pt(13), TEAL
        if gif is not None and os.path.isfile(gif):
            slide.shapes.add_picture(gif, left, Inches(1.3), width=col_width)
        else:
            missing_box = slide.shapes.add_textbox(left, Inches(2.5), col_width, Inches(0.4))
            missing_box.text_frame.text = '(3D render not available for this scan)'
            mrun = missing_box.text_frame.paragraphs[0].runs[0]
            mrun.font.name, mrun.font.size, mrun.font.color.rgb = 'Segoe UI', Pt(12), DARK
    return slide

#----------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--pptx', required=True)
    p.add_argument('--old-analysis-dir', required=True)
    p.add_argument('--new-analysis-dir', required=True)
    p.add_argument('--old-label', required=True)
    p.add_argument('--new-label', required=True)
    p.add_argument('--out', required=True)
    p.add_argument('--max-gif-bytes', type=int, default=3 * 1024 * 1024)
    args = p.parse_args()

    old_viz_index = read_json(os.path.join(args.old_analysis_dir, 'viz_index.json'))
    new_viz_index = read_json(os.path.join(args.new_analysis_dir, 'viz_index.json'))
    scans = shared_scans(old_viz_index, new_viz_index)
    if not scans:
        raise SystemExit(
            'no shared scans between --old-analysis-dir and --new-analysis-dir viz_index.json '
            '(same check build_comparison_deck.py itself would fail on)')

    prs = Presentation(args.pptx)
    layout = find_blank_layout(prs)

    expected = {eyebrow_for(s): s for s in scans}
    slide_index_for_title = {}
    for i, slide in enumerate(prs.slides):
        title = slide_title_text(slide)
        if title in expected and title not in slide_index_for_title:
            slide_index_for_title[title] = i
    missing = set(expected) - set(slide_index_for_title)
    if missing:
        raise SystemExit(
            f'{len(missing)} shared-scan slide(s) not found in --pptx by exact title match: '
            f'{sorted(missing)}. Was --pptx built from the same --old-analysis-dir/'
            f'--new-analysis-dir (or their viz_index.json since changed)?')

    tmp_dir = tempfile.mkdtemp(prefix='augment_deck_3d_')
    n_inserted = 0
    n_missing_gifs = 0
    try:
        # Process from the LAST scan slide to the FIRST: appending+moving a
        # slide to a position after idx never disturbs the recorded index of
        # an earlier, not-yet-processed original scan slide.
        ordered = sorted(slide_index_for_title.items(), key=lambda kv: kv[1], reverse=True)
        for title, slide_idx in ordered:
            scan = expected[title]
            old_gif_src = gif_path(args.old_analysis_dir, scan.patient_id, scan.scan_id)
            new_gif_src = gif_path(args.new_analysis_dir, scan.patient_id, scan.scan_id)
            old_gif = None
            new_gif = None
            if os.path.isfile(old_gif_src):
                old_gif = downsample_gif_under(
                    old_gif_src, os.path.join(tmp_dir, f'old_{scan.patient_id}_{scan.scan_id}.gif'),
                    args.max_gif_bytes)
            if os.path.isfile(new_gif_src):
                new_gif = downsample_gif_under(
                    new_gif_src, os.path.join(tmp_dir, f'new_{scan.patient_id}_{scan.scan_id}.gif'),
                    args.max_gif_bytes)
            if old_gif is None:
                print(f'[warn] no old/flow render_3d.gif for patient={scan.patient_id} scan={scan.scan_id} '
                      f'({old_gif_src!r})')
            if new_gif is None:
                print(f'[warn] no new/EDM2 render_3d.gif for patient={scan.patient_id} scan={scan.scan_id} '
                      f'({new_gif_src!r})')
            if old_gif is None and new_gif is None:
                n_missing_gifs += 1
                continue

            add_companion_slide(prs, layout, scan, old_gif, new_gif, args.old_label, args.new_label)
            new_slide_position = len(prs.slides) - 1
            move_slide(prs, new_slide_position, slide_idx + 1)
            n_inserted += 1

        os.makedirs(os.path.dirname(os.path.abspath(args.out)) or '.', exist_ok=True)
        prs.save(args.out)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    print(f'Wrote {args.out}: {n_inserted} companion slide(s) inserted '
          f'({n_missing_gifs} scan(s) skipped, no GIFs on either side); '
          f'total slides: {len(prs.slides)}')


if __name__ == '__main__':
    main()
